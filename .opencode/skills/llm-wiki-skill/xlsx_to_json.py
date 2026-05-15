#!/usr/bin/env python3
"""Convert XLSX workbooks into a comprehensive JSON representation."""

from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell


def _json_safe(value: Any) -> Any:
    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value.isoformat()
        return value.astimezone(timezone.utc).isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _serialize_cell(cell: Cell) -> dict[str, Any]:
    value = cell.value
    return {
        "address": cell.coordinate,
        "row": cell.row,
        "column": cell.column,
        "column_letter": cell.column_letter,
        "value": _json_safe(value),
        "value_type": type(value).__name__ if value is not None else "NoneType",
        "formula": value if isinstance(value, str) and value.startswith("=") else None,
        "data_type": cell.data_type,
        "number_format": cell.number_format,
        "is_date": bool(cell.is_date),
        "comment": cell.comment.text if cell.comment else None,
        "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
    }


def workbook_to_json(input_path: Path) -> dict[str, Any]:
    wb = load_workbook(filename=input_path, data_only=False)
    props = wb.properties

    document_properties = {
        "title": props.title,
        "subject": props.subject,
        "creator": props.creator,
        "keywords": props.keywords,
        "description": props.description,
        "last_modified_by": props.lastModifiedBy,
        "created": _json_safe(props.created),
        "modified": _json_safe(props.modified),
        "category": props.category,
    }

    workbook_data: dict[str, Any] = {
        "source_file": str(input_path),
        "extracted_at": datetime.now(timezone.utc).isoformat(),
        "workbook": {
            "sheet_names": wb.sheetnames,
            "active_sheet": wb.active.title if wb.active else None,
            "properties": document_properties,
            "defined_names": [
                {
                    "name": defined_name.name,
                    "value": defined_name.attr_text,
                }
                for defined_name in wb.defined_names.values()
            ],
        },
        "sheets": [],
    }

    for sheet in wb.worksheets:
        merged_ranges = [str(cell_range) for cell_range in sheet.merged_cells.ranges]
        table_names = list(sheet.tables.keys())
        cells: list[dict[str, Any]] = []

        for row in sheet.iter_rows(
            min_row=sheet.min_row,
            max_row=sheet.max_row,
            min_col=sheet.min_column,
            max_col=sheet.max_column,
        ):
            for cell in row:
                if cell.value is None and cell.comment is None and cell.hyperlink is None:
                    continue
                cells.append(_serialize_cell(cell))

        workbook_data["sheets"].append(
            {
                "name": sheet.title,
                "dimensions": {
                    "min_row": sheet.min_row,
                    "max_row": sheet.max_row,
                    "min_column": sheet.min_column,
                    "max_column": sheet.max_column,
                },
                "view": {
                    "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                    "auto_filter": str(sheet.auto_filter.ref) if sheet.auto_filter else None,
                },
                "merged_cells": merged_ranges,
                "tables": table_names,
                "cells": cells,
            }
        )

    return workbook_data


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract workbook metadata and all populated cells from XLSX to JSON."
    )
    parser.add_argument("input", type=Path, help="Path to input .xlsx file")
    parser.add_argument(
        "output",
        type=Path,
        nargs="?",
        help="Output JSON path (default: raw/<input-stem>.json)",
    )
    parser.add_argument(
        "--indent",
        type=int,
        default=2,
        help="Pretty print indentation (default: 2)",
    )
    args = parser.parse_args()

    input_path: Path = args.input
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    if input_path.suffix.lower() != ".xlsx":
        raise ValueError(f"Only .xlsx files are supported, got: {input_path.suffix}")

    output_path = args.output or Path("raw") / f"{input_path.stem}.json"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    data = workbook_to_json(input_path)
    with output_path.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=args.indent)
        f.write("\n")

    print(f"Wrote JSON extract to: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
