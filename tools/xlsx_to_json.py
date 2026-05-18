#!/usr/bin/env python3
"""Extract XLSX workbooks to structured JSON with high-fidelity cell metadata."""

from __future__ import annotations

import argparse
import base64
import json
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell


def _serialize_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).isoformat() if value.tzinfo else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bytes):
        return {"encoding": "base64", "data": base64.b64encode(value).decode("ascii")}
    return value


def _value_type(value: Any) -> str:
    return type(value).__name__ if value is not None else "NoneType"


def _serialize_cell(cell: Cell, cached_value: Any) -> dict[str, Any]:
    formula = cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None
    return {
        "address": cell.coordinate,
        "row": cell.row,
        "column": cell.column,
        "column_letter": cell.column_letter,
        "value": _serialize_value(cell.value),
        "value_type": _value_type(cell.value),
        "cached_value": _serialize_value(cached_value),
        "cached_value_type": _value_type(cached_value),
        "formula": formula,
        "data_type": cell.data_type,
        "number_format": cell.number_format,
        "is_date": bool(cell.is_date),
        "style_id": cell.style_id,
        "comment": cell.comment.text if cell.comment else None,
        "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
    }


def workbook_to_json(input_path: Path, include_empty_rows: bool = False) -> dict[str, Any]:
    workbook = load_workbook(filename=input_path, data_only=False, read_only=False)
    workbook_data_only = load_workbook(filename=input_path, data_only=True, read_only=False)
    props = workbook.properties

    output: dict[str, Any] = {
        "source_file": str(input_path),
        "extracted_at": datetime.now(timezone.utc).isoformat(),
        "workbook": {
            "sheet_names": workbook.sheetnames,
            "active_sheet": workbook.active.title if workbook.active else None,
            "properties": {
                "title": props.title,
                "subject": props.subject,
                "creator": props.creator,
                "keywords": props.keywords,
                "description": props.description,
                "last_modified_by": props.lastModifiedBy,
                "created": _serialize_value(props.created),
                "modified": _serialize_value(props.modified),
                "category": props.category,
            },
            "defined_names": [
                {"name": item.name, "value": item.attr_text}
                for item in workbook.defined_names.values()
            ],
        },
        "sheets": [],
    }

    for sheet in workbook.worksheets:
        data_sheet = workbook_data_only[sheet.title]

        merged_ranges = [str(cell_range) for cell_range in sheet.merged_cells.ranges]
        table_names = list(sheet.tables.keys())

        populated_coordinates = set(sheet._cells.keys()) | set(data_sheet._cells.keys())
        by_row: dict[int, list[dict[str, Any]]] = defaultdict(list)
        all_cells: list[dict[str, Any]] = []

        for row_idx, col_idx in sorted(populated_coordinates):
            cell = sheet.cell(row=row_idx, column=col_idx)
            data_cell = data_sheet.cell(row=row_idx, column=col_idx)
            if (
                cell.value is None
                and data_cell.value is None
                and cell.comment is None
                and cell.hyperlink is None
            ):
                continue

            serialized = _serialize_cell(cell, data_cell.value)
            by_row[row_idx].append(serialized)
            all_cells.append(serialized)

        if all_cells:
            rows_present = sorted(by_row.keys())
            min_row = rows_present[0]
            max_row = rows_present[-1]
            columns_present = [cell["column"] for cell in all_cells]
            min_col = min(columns_present)
            max_col = max(columns_present)
        else:
            min_row = 1
            max_row = 1
            min_col = 1
            max_col = 1

        row_records: list[dict[str, Any]] = []
        if include_empty_rows:
            for row_idx in range(min_row, max_row + 1):
                row_cells = by_row.get(row_idx, [])
                row_records.append(
                    {
                        "row": row_idx,
                        "non_empty_cells": len(row_cells),
                        "cells": row_cells,
                    }
                )
        else:
            for row_idx in sorted(by_row.keys()):
                row_cells = by_row[row_idx]
                row_records.append(
                    {
                        "row": row_idx,
                        "non_empty_cells": len(row_cells),
                        "cells": row_cells,
                    }
                )

        output["sheets"].append(
            {
                "name": sheet.title,
                "dimensions": {
                    "min_row": min_row,
                    "max_row": max_row,
                    "min_column": min_col,
                    "max_column": max_col,
                },
                "view": {
                    "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                    "auto_filter": str(sheet.auto_filter.ref) if sheet.auto_filter else None,
                },
                "merged_cells": merged_ranges,
                "tables": table_names,
                "row_count": len(row_records),
                "populated_cell_count": len(all_cells),
                "rows": row_records,
            }
        )

    workbook.close()
    workbook_data_only.close()
    return output


def _default_output_path(input_file: Path, output_arg: Path | None, output_dir: Path | None) -> Path:
    if output_arg is not None and output_dir is not None:
        raise ValueError("Use either explicit output path or --output-dir, not both.")
    if output_arg is not None:
        return output_arg
    if output_dir is not None:
        return output_dir / f"{input_file.stem}.json"
    return input_file.with_suffix(".json")


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract XLSX workbook content to JSON.")
    parser.add_argument("input", type=Path, help="Path to input .xlsx file")
    parser.add_argument("output", type=Path, nargs="?", help="Optional output .json file path")
    parser.add_argument(
        "--output-dir",
        type=Path,
        help="Write JSON into this directory using <input-stem>.json",
    )
    parser.add_argument(
        "--indent",
        type=int,
        default=2,
        help="Pretty-print indentation for JSON output (default: 2)",
    )
    parser.add_argument(
        "--include-empty-rows",
        action="store_true",
        help="Include rows with no populated cells",
    )
    parser.add_argument(
        "--compact",
        action="store_true",
        help="Write compact JSON without extra whitespace",
    )
    args = parser.parse_args()

    input_path = args.input
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError(f"Only .xlsx files are supported, got: {input_path.suffix}")

    output_path = _default_output_path(input_path, args.output, args.output_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    data = workbook_to_json(input_path=input_path, include_empty_rows=args.include_empty_rows)
    with output_path.open("w", encoding="utf-8") as f:
        if args.compact:
            json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
        else:
            json.dump(data, f, ensure_ascii=False, indent=args.indent)
        f.write("\n")

    print(f"Wrote JSON extract to: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
